import io

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _write_sheet(ws: Worksheet, headers: list[str], rows: list[list]) -> None:
    ws.append(headers)
    for row in rows:
        ws.append(row)
    # openpyxl stores any text starting with "=" as a formula. Customer names and messages come from the public chat,
    # so a visitor calling themselves '=HYPERLINK("http://evil…","Refund")' would plant a live formula in the owner's
    # report. Every value we write is data: keep it as text.
    for cells in ws.iter_rows():
        for cell in cells:
            if cell.data_type == "f":
                cell.data_type = "s"


def build_report_workbook(report: dict) -> Workbook:
    """Real .xlsx via openpyxl — same report dict generate_daily_report already
    produces, so the JSON endpoint and this export can never disagree.
    Sheet order matches the master plan: Appointments, Cancellations,
    Reschedules, New Leads, Summary."""
    wb = Workbook()

    ws = wb.active
    ws.title = "Appointments"
    _write_sheet(
        ws,
        ["Time", "Customer", "Service", "Staff", "Status", "Booking ID"],
        [
            [a["scheduled_at"], a["customer_name"], a["service_name"], a["staff_name"], a["status"], a["id"]]
            for a in report["appointments"]
        ],
    )

    ws = wb.create_sheet("Cancellations")
    _write_sheet(
        ws,
        ["Originally Scheduled", "Cancelled At", "Customer", "Service", "Booking ID"],
        [
            [c["originally_scheduled_at"], c["cancelled_at"], c["customer_name"], c["service_name"], c["id"]]
            for c in report["cancellations"]
        ],
    )

    ws = wb.create_sheet("Reschedules")
    _write_sheet(
        ws,
        ["Old Time", "New Time", "Changed At", "Customer", "Service", "Booking ID"],
        [
            [r["old_scheduled_at"], r["new_scheduled_at"], r["changed_at"], r["customer_name"], r["service_name"], r["id"]]
            for r in report["reschedules"]
        ],
    )

    ws = wb.create_sheet("New Leads")
    _write_sheet(
        ws,
        ["Name", "Phone", "Email", "Created At", "Customer ID"],
        [[c["name"], c["phone"], c["email"], c["created_at"], c["id"]] for c in report["new_leads"]],
    )

    ws = wb.create_sheet("Summary")
    summary = report["summary"]
    rows = [
        ["Business", report["business_name"]],
        ["Report Date", report["report_date"]],
        ["Timezone", report["timezone"]],
        ["Appointments Scheduled", summary["appointments_scheduled"]],
        *[[f"  — {status}", count] for status, count in summary["appointments_by_status"].items()],
        ["Cancellations", summary["cancellations"]],
        ["Reschedules", summary["reschedules"]],
        ["New Leads", summary["new_leads"]],
        ["Human Review — Open Count", summary["human_review_open_count"]],
        ["Human Review — Feature Implemented", report["human_review"]["implemented"]],
        ["Human Review — Note", report["human_review"]["note"]],
        ["Revenue Estimate", report["revenue_estimate"]["value"]],
        ["Revenue Estimate — Definition", report["revenue_estimate"]["definition"]],
    ]
    _write_sheet(ws, ["Metric", "Value"], rows)

    return wb


def report_to_xlsx_bytes(report: dict) -> bytes:
    wb = build_report_workbook(report)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_monthly_report_workbook(report: dict) -> Workbook:
    """Same pattern as build_report_workbook — real openpyxl workbook from the
    exact dict generate_monthly_report produces, so the JSON endpoint and this
    export can never disagree. Sheets: Summary, Busiest Days, Busiest Hours,
    Most Requested Services."""
    wb = Workbook()
    a = report["appointments"]
    cr = report["cancellation_rate"]
    bc = report["booking_conversion"]

    ws = wb.active
    ws.title = "Summary"
    _write_sheet(
        ws,
        ["Metric", "Value"],
        [
            ["Business", report["business_name"]],
            ["Period", report["period_label"]],
            ["Timezone", report["timezone"]],
            ["Conversations", report["conversations"]["total"]],
            ["New Customers", report["customers"]["new"]],
            ["Total Customers (month end)", report["customers"]["total_at_month_end"]],
            ["Appointments Requested", a["requested"]],
            ["Appointments Scheduled This Month", a["scheduled_for_month"]],
            ["Cancellation Events This Month", a["cancellation_events_this_month"]],
            ["Cancelled (of scheduled-for-month)", a["cancelled_of_scheduled"]],
            ["Cancellation Rate", cr["value"]],
            ["Cancellation Rate — Definition", cr["definition"]],
            ["Booking Conversion", bc["value"]],
            ["Booking Conversion — Definition", bc["definition"]],
            ["Reschedule Events", a["rescheduled"]["events"]],
            ["Reschedule — Distinct Appointments", a["rescheduled"]["distinct_appointments"]],
            ["Completed (real count)", a["completed"]["count"]],
            ["Completed — Feature Implemented", a["completed"]["implemented"]],
            ["Completed — Note", a["completed"]["note"]],
            ["Revenue Estimate", report["revenue_estimate"]["value"]],
            ["Revenue Estimate — Definition", report["revenue_estimate"]["definition"]],
        ],
    )

    ws = wb.create_sheet("Busiest Days")
    _write_sheet(ws, ["Day", "Count"], [[d["day"], d["count"]] for d in report["busiest_days"]])

    ws = wb.create_sheet("Busiest Hours")
    _write_sheet(ws, ["Hour", "Count"], [[h["hour"], h["count"]] for h in report["busiest_hours"]])

    ws = wb.create_sheet("Most Requested Services")
    _write_sheet(
        ws,
        ["Service", "Count", "Service ID"],
        [[s["service_name"], s["count"], s["service_id"]] for s in report["most_requested_services"]],
    )

    return wb


def monthly_report_to_xlsx_bytes(report: dict) -> bytes:
    wb = build_monthly_report_workbook(report)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_yearly_report_workbook(report: dict) -> Workbook:
    """Same pattern as build_monthly_report_workbook — real openpyxl workbook
    from the exact dict generate_yearly_report produces. Sheets: Summary,
    Month by Month, Most Requested Services."""
    wb = Workbook()
    a = report["appointments"]
    cr = report["cancellation_rate"]
    bc = report["booking_conversion"]
    rev = report["revenue_estimate"]

    ws = wb.active
    ws.title = "Summary"
    rows = [
        ["Business", report["business_name"]],
        ["Year", report["year"]],
        ["Timezone", report["timezone"]],
        ["Conversations", report["conversations"]["total"]],
        ["New Customers", report["customers"]["new"]],
        ["Appointments Requested", a["requested"]],
        ["Appointments Scheduled This Year", a["scheduled_for_year"]],
        ["Cancelled (of scheduled-for-year)", a["cancelled_of_scheduled"]],
        ["Cancellation Rate", cr["value"]],
        ["Cancellation Rate — Definition", cr["definition"]],
        ["Booking Conversion", bc["value"]],
        ["Booking Conversion — Definition", bc["definition"]],
        ["Reschedule Events", a["rescheduled"]["events"]],
        ["Completed (real count)", a["completed"]["count"]],
        ["Completed — Feature Implemented", a["completed"]["implemented"]],
        ["Revenue Estimate", rev["value"]],
        ["Revenue Estimate — Definition", rev["definition"]],
    ]
    yoy = report.get("year_over_year")
    if yoy:
        rows.append(["Year-over-Year Available", yoy["available"]])
        if yoy["available"]:
            rows += [
                ["YoY Prior Year", yoy["prior_year"]],
                ["YoY Appointments Scheduled (current)", yoy["appointments_scheduled"]["current"]],
                ["YoY Appointments Scheduled (prior)", yoy["appointments_scheduled"]["prior"]],
                ["YoY Appointments Change %", yoy["appointments_scheduled"]["change_pct"]],
                ["YoY Revenue Estimate (current)", yoy["revenue_estimate"]["current"]],
                ["YoY Revenue Estimate (prior)", yoy["revenue_estimate"]["prior"]],
                ["YoY Revenue Change %", yoy["revenue_estimate"]["change_pct"]],
                ["YoY New Customers (current)", yoy["new_customers"]["current"]],
                ["YoY New Customers (prior)", yoy["new_customers"]["prior"]],
                ["YoY New Customers Change %", yoy["new_customers"]["change_pct"]],
            ]
        else:
            rows += [["YoY Prior Year", yoy["prior_year"]], ["Year-over-Year Note", yoy["note"]]]
    _write_sheet(ws, ["Metric", "Value"], rows)

    ws = wb.create_sheet("Month by Month")
    _write_sheet(
        ws,
        ["Month", "Requested", "Scheduled", "Cancelled", "Revenue Estimate"],
        [
            [m["month_name"], m["appointments_requested"], m["appointments_scheduled"], m["cancelled"], m["revenue_estimate"]]
            for m in report["month_by_month"]
        ],
    )

    ws = wb.create_sheet("Most Requested Services")
    _write_sheet(
        ws,
        ["Service", "Count", "Service ID"],
        [[s["service_name"], s["count"], s["service_id"]] for s in report["most_requested_services"]],
    )

    return wb


def yearly_report_to_xlsx_bytes(report: dict) -> bytes:
    wb = build_yearly_report_workbook(report)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
