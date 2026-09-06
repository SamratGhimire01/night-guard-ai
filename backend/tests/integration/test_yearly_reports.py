"""Phase 37 — the yearly report (app/services/reporting/yearly_report_service.py,
GET /api/v1/reports/yearly*), gated to BusinessPlan.PREMIUM via require_plan
(Phase 34) — the first REAL premium-gated feature (Phase 34's premium_test.py
was a synthetic scaffold proving the mechanism worked; this is the genuine
article).

Real DB and real HTTP throughout for the booking/customer/conversation data
(never pre-shaped fixtures) except where the live API structurally cannot
produce the data needed (a real appointment in a PAST calendar year, for the
year-over-year test) — booking_service rejects past-dated appointments by
design, so that one test follows the exact same "direct ORM Appointment(...)
insert, call the report function directly" precedent already established in
test_monthly_reports.py's own test_month_boundary_uses_business_local_timezone_not_utc.
"""

import calendar
import uuid
from datetime import date, datetime, timedelta
from decimal import Decimal
from zoneinfo import ZoneInfo

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.core.security import create_access_token, hash_password
from app.db.database import SessionLocal
from app.db.models.appointment import Appointment, AppointmentStatus
from app.db.models.business import Business, BusinessPlan, BusinessUser, BusinessUserRole
from app.db.models.customer import Customer
from app.db.models.service import Service
from app.main import app
from app.services.reporting import yearly_report_service

client = TestClient(app)


def _unique_email(label: str) -> str:
    return f"{label}-{uuid.uuid4().hex[:10]}@example.com"


def _auth_header(token: str) -> dict:
    return {"Authorization": f"Bearer {token}"}


def _day_in_current_month(offset: int) -> date:
    today = date.today()
    candidate = today + timedelta(days=offset)
    if candidate.month != today.month or candidate.year != today.year:
        last_day = calendar.monthrange(today.year, today.month)[1]
        candidate = today.replace(day=last_day)
    return candidate


@pytest.fixture
def two_businesses():
    email_a = _unique_email("yrpt-a-owner")
    email_b = _unique_email("yrpt-b-owner")
    resp_a = client.post(
        "/api/v1/auth/register",
        json={"business_name": "Yearly Report A", "timezone": "UTC", "email": email_a, "password": "correcthorse1"},
    )
    resp_b = client.post(
        "/api/v1/auth/register",
        json={"business_name": "Yearly Report B", "timezone": "UTC", "email": email_b, "password": "correcthorse1"},
    )
    assert resp_a.status_code == 201, resp_a.text
    assert resp_b.status_code == 201, resp_b.text
    login_a = client.post("/api/v1/auth/login", json={"email": email_a, "password": "correcthorse1"})
    login_b = client.post("/api/v1/auth/login", json={"email": email_b, "password": "correcthorse1"})
    data = {
        "business_id_a": resp_a.json()["business_id"],
        "business_id_b": resp_b.json()["business_id"],
        "token_a": login_a.json()["access_token"],
        "token_b": login_b.json()["access_token"],
    }
    yield data
    with SessionLocal() as db:
        for business_id in (data["business_id_a"], data["business_id_b"]):
            business = db.get(Business, uuid.UUID(business_id))
            if business is not None:
                db.delete(business)
        db.commit()


@pytest.fixture
def staff_token(two_businesses):
    business_id_a = uuid.UUID(two_businesses["business_id_a"])
    with SessionLocal() as db:
        staff_user = BusinessUser(
            business_id=business_id_a,
            email=_unique_email("yrpt-staff"),
            hashed_password="unused",
            role=BusinessUserRole.STAFF,
        )
        db.add(staff_user)
        db.commit()
        db.refresh(staff_user)
        return create_access_token(user_id=staff_user.id, business_id=business_id_a, role=staff_user.role.value)


@pytest.fixture
def superadmin_token():
    """Same technique test_plan_entitlement.py's own fixture uses — no API
    exists to grant is_superadmin (a deliberate Phase 34 gap), so this mints
    one via direct ORM insert."""
    with SessionLocal() as db:
        biz = Business(name="Yearly Test Ops (internal)", timezone="UTC")
        db.add(biz)
        db.flush()
        admin_user = BusinessUser(
            business_id=biz.id,
            email=_unique_email("yrpt-platform-admin"),
            hashed_password=hash_password("not-used-in-this-test"),
            role=BusinessUserRole.OWNER,
            is_superadmin=True,
        )
        db.add(admin_user)
        db.commit()
        db.refresh(admin_user)
        token = create_access_token(
            user_id=admin_user.id, business_id=admin_user.business_id, role=admin_user.role.value
        )
        business_id = biz.id

    yield token

    with SessionLocal() as db:
        business = db.get(Business, business_id)
        if business is not None:
            db.delete(business)
        db.commit()


def _open_all_week(token: str):
    days = [{"day_of_week": d, "open_time": "00:00:00", "close_time": "23:45:00"} for d in range(7)]
    assert client.put("/api/v1/business/hours", json={"days": days}, headers=_auth_header(token)).status_code == 200


def _create_service(token: str, name: str = "Cleaning", price: str = "50.00") -> uuid.UUID:
    resp = client.post(
        "/api/v1/services", json={"name": name, "price": price, "duration_minutes": 30}, headers=_auth_header(token)
    )
    assert resp.status_code == 201, resp.text
    return uuid.UUID(resp.json()["id"])


def _create_customer(token: str, phone: str) -> dict:
    resp = client.post(
        "/api/v1/customers", json={"name": "Yearly Customer", "phone": phone}, headers=_auth_header(token)
    )
    assert resp.status_code == 201, resp.text
    return resp.json()


def _book(token: str, service_id: uuid.UUID, customer_id: uuid.UUID, when: datetime) -> dict:
    resp = client.post(
        "/api/v1/appointments",
        json={"customer_id": str(customer_id), "service_id": str(service_id), "scheduled_at": when.isoformat()},
        headers=_auth_header(token),
    )
    assert resp.status_code == 201, resp.text
    return resp.json()


def _upgrade_to_premium(superadmin_token: str, business_id: str) -> dict:
    resp = client.patch(
        f"/api/v1/admin/businesses/{business_id}/plan",
        json={"plan": "premium"},
        headers=_auth_header(superadmin_token),
    )
    assert resp.status_code == 200, resp.text
    return resp.json()


@pytest.fixture
def business_ready(two_businesses):
    token = two_businesses["token_a"]
    _open_all_week(token)
    service_id = _create_service(token)
    return {"token": token, "business_id": uuid.UUID(two_businesses["business_id_a"]), "service_id": service_id}


def _this_year() -> int:
    return date.today().year


# --- plan gating: the real, non-scaffold premium gate --------------------------------------


def test_yearly_report_returns_402_for_free_plan(business_ready):
    resp = client.get(f"/api/v1/reports/yearly?year={_this_year()}", headers=_auth_header(business_ready["token"]))
    assert resp.status_code == 402, resp.text
    assert resp.json()["error"]["type"] == "plan_required"


def test_yearly_report_renders_after_real_premium_upgrade(business_ready, two_businesses, superadmin_token):
    token = business_ready["token"]
    year = _this_year()

    blocked = client.get(f"/api/v1/reports/yearly?year={year}", headers=_auth_header(token))
    assert blocked.status_code == 402

    _upgrade_to_premium(superadmin_token, two_businesses["business_id_a"])

    allowed = client.get(f"/api/v1/reports/yearly?year={year}", headers=_auth_header(token))
    assert allowed.status_code == 200, allowed.text
    report = allowed.json()
    assert report["year"] == year
    assert report["business_name"] == "Yearly Report A"


def test_yearly_report_staff_forbidden_regardless_of_plan(
    business_ready, two_businesses, staff_token, superadmin_token
):
    _upgrade_to_premium(superadmin_token, two_businesses["business_id_a"])
    resp = client.get(f"/api/v1/reports/yearly?year={_this_year()}", headers=_auth_header(staff_token))
    assert resp.status_code == 403, resp.text


# --- real aggregation matches real DB state -------------------------------------------------


def test_yearly_report_matches_real_db_after_upgrade(business_ready, two_businesses, superadmin_token):
    token, service_id = business_ready["token"], business_ready["service_id"]
    _upgrade_to_premium(superadmin_token, two_businesses["business_id_a"])
    year = _this_year()

    customer1 = _create_customer(token, phone="+15556660001")
    customer2 = _create_customer(token, phone="+15556660002")
    when1 = datetime.combine(_day_in_current_month(1), datetime.min.time()).replace(hour=9, tzinfo=ZoneInfo("UTC"))
    when2 = datetime.combine(_day_in_current_month(2), datetime.min.time()).replace(hour=10, tzinfo=ZoneInfo("UTC"))
    appt1 = _book(token, service_id, uuid.UUID(customer1["id"]), when1)
    _book(token, service_id, uuid.UUID(customer2["id"]), when2)
    client.patch(f"/api/v1/appointments/{appt1['id']}/cancel", headers=_auth_header(token))

    report = client.get(f"/api/v1/reports/yearly?year={year}", headers=_auth_header(token)).json()

    with SessionLocal() as db:
        real_total = (
            db.query(Appointment).filter(Appointment.business_id == business_ready["business_id"]).count()
        )
        real_cancelled = (
            db.query(Appointment)
            .filter(
                Appointment.business_id == business_ready["business_id"],
                Appointment.status == AppointmentStatus.CANCELLED,
            )
            .count()
        )

    assert report["appointments"]["requested"] == real_total == 2
    assert report["appointments"]["scheduled_for_year"] == 2
    assert report["appointments"]["cancelled_of_scheduled"] == real_cancelled == 1
    # Revenue: only the non-cancelled appointment (customer2) counts, at the
    # service's real $50.00 price.
    assert report["revenue_estimate"]["value"] == "50.00"
    assert report["revenue_estimate"]["appointment_count"] == 1
    # month_by_month sums back to the same yearly totals (single source of
    # truth check: yearly is built FROM these 12 monthly rows, so it must).
    assert sum(m["appointments_scheduled"] for m in report["month_by_month"]) == 2
    assert sum(m["cancelled"] for m in report["month_by_month"]) == 1
    assert sum(Decimal(m["revenue_estimate"]) for m in report["month_by_month"]) == Decimal("50.00")


def test_yearly_report_zero_activity_year_is_honest(business_ready, two_businesses, superadmin_token):
    _upgrade_to_premium(superadmin_token, two_businesses["business_id_a"])
    resp = client.get("/api/v1/reports/yearly?year=2015", headers=_auth_header(business_ready["token"]))
    assert resp.status_code == 200, resp.text
    report = resp.json()
    assert report["appointments"]["requested"] == 0
    assert report["appointments"]["scheduled_for_year"] == 0
    assert report["cancellation_rate"]["value"] is None
    assert report["booking_conversion"]["value"] is None
    assert report["revenue_estimate"]["value"] == "0.00"
    assert all(m["appointments_scheduled"] == 0 for m in report["month_by_month"])
    assert report["most_requested_services"] == []
    # 2014 also has zero real activity — no fabricated year-over-year comparison.
    assert report["year_over_year"]["available"] is False
    assert report["year_over_year"]["prior_year"] == 2014


# --- cross-tenant isolation ------------------------------------------------------------------


def test_yearly_report_cross_tenant_isolation(business_ready, two_businesses, superadmin_token):
    token_a, token_b = business_ready["token"], two_businesses["token_b"]
    _upgrade_to_premium(superadmin_token, two_businesses["business_id_a"])
    _upgrade_to_premium(superadmin_token, two_businesses["business_id_b"])
    year = _this_year()

    _open_all_week(token_b)
    service_b = _create_service(token_b, name="B Service", price="75.00")
    customer_b = _create_customer(token_b, phone="+15557770001")
    when = datetime.combine(_day_in_current_month(1), datetime.min.time()).replace(hour=14, tzinfo=ZoneInfo("UTC"))
    appt_b = _book(token_b, service_b, uuid.UUID(customer_b["id"]), when)

    report_a = client.get(f"/api/v1/reports/yearly?year={year}", headers=_auth_header(token_a)).json()
    report_b = client.get(f"/api/v1/reports/yearly?year={year}", headers=_auth_header(token_b)).json()

    assert report_a["appointments"]["requested"] == 0
    assert report_b["appointments"]["requested"] == 1
    assert report_a["business_name"] == "Yearly Report A"
    assert report_b["business_name"] == "Yearly Report B"
    assert appt_b["id"] not in str(report_a)


# --- year-over-year: real comparison when prior-year data genuinely exists -------------------


def test_year_over_year_comparison_with_real_prior_year_data():
    """The live booking API structurally cannot create a PAST-dated
    appointment (booking_service rejects past times by design), so this
    follows test_monthly_reports.py's own precedent for the identical
    problem: direct ORM Appointment(...) rows, then call the report
    function directly (not via HTTP) — same real aggregation code path
    GET /reports/yearly uses, just invoked without an HTTP round trip."""
    year = 2020
    with SessionLocal() as db:
        business = Business(name="YoY Real Data Co", timezone="UTC", plan=BusinessPlan.PREMIUM)
        db.add(business)
        db.flush()
        service = Service(business_id=business.id, name="Checkup", price=Decimal("100.00"), duration_minutes=30)
        customer = Customer(business_id=business.id, name="YoY Customer")
        db.add_all([service, customer])
        db.flush()

        # Prior year (2019): 2 real non-cancelled appointments -> $200 revenue.
        for month, day in [(3, 10), (6, 15)]:
            db.add(
                Appointment(
                    business_id=business.id,
                    customer_id=customer.id,
                    service_id=service.id,
                    scheduled_at=datetime(year - 1, month, day, 9, tzinfo=ZoneInfo("UTC")),
                    duration_minutes=30,
                    status=AppointmentStatus.CONFIRMED,
                )
            )
        # Current year (2020): 4 real non-cancelled appointments -> $400 revenue.
        for month, day in [(1, 5), (2, 5), (3, 5), (4, 5)]:
            db.add(
                Appointment(
                    business_id=business.id,
                    customer_id=customer.id,
                    service_id=service.id,
                    scheduled_at=datetime(year, month, day, 9, tzinfo=ZoneInfo("UTC")),
                    duration_minutes=30,
                    status=AppointmentStatus.CONFIRMED,
                )
            )
        db.commit()
        business_id = business.id

    with SessionLocal() as db:
        report = yearly_report_service.generate_yearly_report(db, business_id=business_id, year=year)

    yoy = report["year_over_year"]
    assert yoy["available"] is True
    assert yoy["prior_year"] == year - 1
    assert yoy["appointments_scheduled"]["current"] == 4
    assert yoy["appointments_scheduled"]["prior"] == 2
    assert yoy["appointments_scheduled"]["change_pct"] == pytest.approx(1.0)  # +100%
    assert yoy["revenue_estimate"]["current"] == "400.00"
    assert yoy["revenue_estimate"]["prior"] == "200.00"
    assert yoy["revenue_estimate"]["change_pct"] == pytest.approx(1.0)

    with SessionLocal() as db:
        business = db.get(Business, business_id)
        db.delete(business)
        db.commit()


# --- real .xlsx export matches the JSON exactly ----------------------------------------------


def test_yearly_excel_export_matches_json(business_ready, two_businesses, superadmin_token):
    token, service_id = business_ready["token"], business_ready["service_id"]
    _upgrade_to_premium(superadmin_token, two_businesses["business_id_a"])
    year = _this_year()

    customer = _create_customer(token, phone="+15558880001")
    when = datetime.combine(_day_in_current_month(1), datetime.min.time()).replace(hour=13, tzinfo=ZoneInfo("UTC"))
    _book(token, service_id, uuid.UUID(customer["id"]), when)

    report = client.get(f"/api/v1/reports/yearly?year={year}", headers=_auth_header(token)).json()

    resp = client.get(f"/api/v1/reports/yearly/excel?year={year}", headers=_auth_header(token))
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert f'yearly_report_{year}.xlsx' in resp.headers["content-disposition"]

    import io

    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["Summary", "Month by Month", "Most Requested Services"]
    summary_rows = {row[0]: row[1] for row in wb["Summary"].iter_rows(values_only=True)}
    assert summary_rows["Appointments Scheduled This Year"] == report["appointments"]["scheduled_for_year"] == 1
    assert summary_rows["Revenue Estimate"] == report["revenue_estimate"]["value"] == "50.00"


# --- revenue_estimate on daily/monthly (Phase 37 additions to existing reports) --------------


def test_daily_and_monthly_revenue_estimate_matches_real_service_price(business_ready):
    token, service_id = business_ready["token"], business_ready["service_id"]
    day = _day_in_current_month(1)

    customer1 = _create_customer(token, phone="+15559990001")
    customer2 = _create_customer(token, phone="+15559990002")
    when1 = datetime.combine(day, datetime.min.time()).replace(hour=9, tzinfo=ZoneInfo("UTC"))
    when2 = datetime.combine(day, datetime.min.time()).replace(hour=11, tzinfo=ZoneInfo("UTC"))
    appt1 = _book(token, service_id, uuid.UUID(customer1["id"]), when1)
    _book(token, service_id, uuid.UUID(customer2["id"]), when2)
    client.patch(f"/api/v1/appointments/{appt1['id']}/cancel", headers=_auth_header(token))

    daily = client.get(f"/api/v1/reports/daily?date={day.isoformat()}", headers=_auth_header(token)).json()
    assert daily["revenue_estimate"]["value"] == "50.00"
    assert daily["revenue_estimate"]["appointment_count"] == 1
    assert "ESTIMATE" in daily["revenue_estimate"]["definition"]
    # revenue_estimate must not have altered the pre-existing exact-equality
    # summary contract from Phase 16 — it lives as a sibling top-level key.
    assert "revenue_estimate" not in daily["summary"]

    monthly = client.get(
        f"/api/v1/reports/monthly?year={day.year}&month={day.month}", headers=_auth_header(token)
    ).json()
    assert monthly["revenue_estimate"]["value"] == "50.00"
    assert monthly["revenue_estimate"]["appointment_count"] == 1
