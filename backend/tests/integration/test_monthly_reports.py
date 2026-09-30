"""Phase 17 Part B — monthly analytics (app/services/reporting/monthly_report_service.py,
GET/POST /api/v1/reports/monthly*).

Real DB and real HTTP throughout for the booking/cancel/reschedule/customer/
conversation data (never pre-shaped fixtures) so the report is proven against
genuinely-produced data. Only the SMTP network call in the email tests is
stubbed.
"""

import calendar
import uuid
from datetime import date, datetime, timedelta
from zoneinfo import ZoneInfo

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.core.security import create_access_token
from app.db.database import SessionLocal
from app.db.models.appointment import Appointment, AppointmentStatus
from app.db.models.business import Business, BusinessUser, BusinessUserRole
from app.db.models.conversation import Conversation
from app.db.models.customer import Customer
from app.db.models.service import Service
from app.main import app
from app.services.reporting import monthly_report_service

client = TestClient(app)


def _unique_email(label: str) -> str:
    return f"{label}-{uuid.uuid4().hex[:10]}@example.com"


def _auth_header(token: str) -> dict:
    return {"Authorization": f"Bearer {token}"}


def _test_month_start() -> date:
    """The month these tests book into and report on: this month while at least 12 future days remain in it,
    otherwise next month. Every booking must be in the future (the booking API refuses the past), and some tests need
    up to 10 distinct days in one month; near the end of a month (e.g. on the 30th) this month can't provide them."""
    today = date.today()
    last_day = calendar.monthrange(today.year, today.month)[1]
    if today.day + 11 <= last_day:
        return today
    return (today.replace(day=1) + timedelta(days=32)).replace(day=1)


def _day_in_current_month(offset: int) -> date:
    """A real future date `offset` days into the test month (see _test_month_start), never spilling out of it."""
    start = _test_month_start()
    candidate = start + timedelta(days=offset)
    if candidate.month != start.month:
        candidate = start.replace(day=calendar.monthrange(start.year, start.month)[1])
    return candidate


@pytest.fixture
def two_businesses():
    email_a = _unique_email("mrpt-a-owner")
    email_b = _unique_email("mrpt-b-owner")
    resp_a = client.post(
        "/api/v1/auth/register",
        json={"business_name": "Monthly Report A", "timezone": "UTC", "email": email_a, "password": "correcthorse1"},
    )
    resp_b = client.post(
        "/api/v1/auth/register",
        json={"business_name": "Monthly Report B", "timezone": "UTC", "email": email_b, "password": "correcthorse1"},
    )
    assert resp_a.status_code == 201, resp_a.text
    assert resp_b.status_code == 201, resp_b.text
    login_a = client.post("/api/v1/auth/login", json={"email": email_a, "password": "correcthorse1"})
    login_b = client.post("/api/v1/auth/login", json={"email": email_b, "password": "correcthorse1"})
    data = {
        "business_id_a": resp_a.json()["business_id"],
        "business_id_b": resp_b.json()["business_id"],
        "token_a": login_a.json()["access_token"],
        "token_b": login_b.json()["access_token"],
    }
    yield data
    with SessionLocal() as db:
        for business_id in (data["business_id_a"], data["business_id_b"]):
            business = db.get(Business, uuid.UUID(business_id))
            if business is not None:
                db.delete(business)
        db.commit()


@pytest.fixture
def staff_token(two_businesses):
    business_id_a = uuid.UUID(two_businesses["business_id_a"])
    with SessionLocal() as db:
        staff_user = BusinessUser(
            business_id=business_id_a,
            email=_unique_email("mrpt-staff"),
            hashed_password="unused",
            role=BusinessUserRole.STAFF,
        )
        db.add(staff_user)
        db.commit()
        db.refresh(staff_user)
        return create_access_token(user_id=staff_user.id, business_id=business_id_a, role=staff_user.role.value)


def _open_all_week(token: str):
    days = [{"day_of_week": d, "open_time": "00:00:00", "close_time": "23:45:00"} for d in range(7)]
    assert client.put("/api/v1/business/hours", json={"days": days}, headers=_auth_header(token)).status_code == 200


def _create_service(token: str, name: str = "Cleaning", duration_minutes: int = 30) -> uuid.UUID:
    resp = client.post(
        "/api/v1/services",
        json={"name": name, "price": "50.00", "duration_minutes": duration_minutes},
        headers=_auth_header(token),
    )
    assert resp.status_code == 201, resp.text
    return uuid.UUID(resp.json()["id"])


def _create_customer(token: str, **kwargs) -> dict:
    payload = {"name": "Monthly Customer", **kwargs}
    resp = client.post("/api/v1/customers", json=payload, headers=_auth_header(token))
    assert resp.status_code == 201, resp.text
    return resp.json()


def _book(token: str, service_id: uuid.UUID, customer_id: uuid.UUID, when: datetime) -> dict:
    resp = client.post(
        "/api/v1/appointments",
        json={"customer_id": str(customer_id), "service_id": str(service_id), "scheduled_at": when.isoformat()},
        headers=_auth_header(token),
    )
    assert resp.status_code == 201, resp.text
    return resp.json()


def _create_conversation(business_id: uuid.UUID, customer_id: uuid.UUID) -> uuid.UUID:
    with SessionLocal() as db:
        conversation = Conversation(business_id=business_id, customer_id=customer_id, channel="sms", status="open")
        db.add(conversation)
        db.commit()
        return conversation.id


@pytest.fixture
def business_ready(two_businesses):
    token = two_businesses["token_a"]
    _open_all_week(token)
    service_id = _create_service(token)
    return {"token": token, "business_id": uuid.UUID(two_businesses["business_id_a"]), "service_id": service_id}


def _this_year_month() -> tuple[int, int]:
    start = _test_month_start()
    return start.year, start.month

def _current_year_month() -> tuple[int, int]:
    """For tests that check what was *created* this month (customers, conversations, requested appointments): those
    rows get today's timestamp, so the report has to be for the current month."""
    today = date.today()
    return today.year, today.month


def _future_slot(offset_days: int, hour: int) -> datetime:
    """A bookable future time `offset_days` from today at `hour` UTC, kept inside the current month. On the month's
    last days that lands today, possibly already past, so it moves to a quarter hour at least 30 minutes ahead, spaced
    by `hour` so bookings never overlap."""
    today = date.today()
    day = today + timedelta(days=offset_days)
    if day.month != today.month:
        day = today
    now = datetime.now(ZoneInfo("UTC"))
    when = datetime.combine(day, datetime.min.time()).replace(hour=hour, tzinfo=ZoneInfo("UTC"))
    if when <= now + timedelta(minutes=30):
        base = now + timedelta(minutes=30 + (hour % 4) * 45)
        when = base.replace(minute=(base.minute // 15) * 15, second=0, microsecond=0) + timedelta(minutes=15)
    return when



# --- real counts cross-checked against raw DB queries -----------------------------------


def test_monthly_report_counts_match_real_db_state(business_ready):
    token, service_id = business_ready["token"], business_ready["service_id"]
    year, month = _current_year_month()

    customer1 = _create_customer(token, phone="+15552220001")
    customer2 = _create_customer(token, phone="+15552220002")
    when1 = _future_slot(1, 9)
    when2 = _future_slot(2, 10)
    appt1 = _book(token, service_id, uuid.UUID(customer1["id"]), when1)
    _appt2 = _book(token, service_id, uuid.UUID(customer2["id"]), when2)
    client.patch(f"/api/v1/appointments/{appt1['id']}/cancel", headers=_auth_header(token))

    _create_conversation(business_ready["business_id"], uuid.UUID(customer1["id"]))
    _create_conversation(business_ready["business_id"], uuid.UUID(customer2["id"]))

    resp = client.get(f"/api/v1/reports/monthly?year={year}&month={month}", headers=_auth_header(token))
    assert resp.status_code == 200, resp.text
    report = resp.json()

    # Manual, independent DB cross-check (not just trusting the endpoint):
    with SessionLocal() as db:
        real_customers = (
            db.query(Customer).filter(Customer.business_id == business_ready["business_id"]).count()
        )
        real_conversations = (
            db.query(Conversation).filter(Conversation.business_id == business_ready["business_id"]).count()
        )
        real_appointments_total = (
            db.query(Appointment).filter(Appointment.business_id == business_ready["business_id"]).count()
        )
        real_cancelled = (
            db.query(Appointment)
            .filter(Appointment.business_id == business_ready["business_id"], Appointment.status == AppointmentStatus.CANCELLED)
            .count()
        )

    assert report["customers"]["new"] == real_customers == 2
    assert report["conversations"]["total"] == real_conversations == 2
    assert report["appointments"]["requested"] == real_appointments_total == 2
    assert report["appointments"]["scheduled_for_month"] == 2
    assert report["appointments"]["cancelled_of_scheduled"] == real_cancelled == 1


def test_zero_activity_month_is_honest(business_ready):
    resp = client.get("/api/v1/reports/monthly?year=2019&month=1", headers=_auth_header(business_ready["token"]))
    assert resp.status_code == 200, resp.text
    report = resp.json()
    assert report["conversations"]["total"] == 0
    assert report["customers"]["new"] == 0
    assert report["appointments"]["requested"] == 0
    assert report["appointments"]["scheduled_for_month"] == 0
    assert report["cancellation_rate"]["value"] is None
    assert report["booking_conversion"]["value"] is None
    assert report["busiest_days"] == [{"day": d, "count": 0} for d in calendar.day_name]
    assert report["busiest_hours"] == []
    assert report["most_requested_services"] == []


# --- booking_conversion: stated definition, verified numerator/denominator --------------


def test_booking_conversion_definition_and_value(business_ready):
    token, service_id = business_ready["token"], business_ready["service_id"]
    year, month = _current_year_month()

    for i in range(4):
        cust = _create_customer(token, phone=f"+1555333000{i}")
        _create_conversation(business_ready["business_id"], uuid.UUID(cust["id"]))
    booking_customer = _create_customer(token, phone="+15553330099")
    when = _future_slot(3, 11)
    _book(token, service_id, uuid.UUID(booking_customer["id"]), when)

    report = client.get(f"/api/v1/reports/monthly?year={year}&month={month}", headers=_auth_header(token)).json()
    bc = report["booking_conversion"]
    assert bc["numerator"] == 1  # 1 appointment requested (created)
    assert bc["denominator"] == 4  # 4 conversations created (the booking customer got none)
    assert bc["value"] == pytest.approx(1 / 4)
    assert "proxy" in bc["definition"].lower()


# --- cancellation_rate: stated definition, verified numerator/denominator ---------------


def test_cancellation_rate_definition_and_value(business_ready):
    token, service_id = business_ready["token"], business_ready["service_id"]
    year, month = _this_year_month()

    appt_ids = []
    for i in range(4):
        cust = _create_customer(token, phone=f"+1555444000{i}")
        when = datetime.combine(_day_in_current_month(4), datetime.min.time()).replace(hour=8 + i, tzinfo=ZoneInfo("UTC"))
        appt = _book(token, service_id, uuid.UUID(cust["id"]), when)
        appt_ids.append(appt["id"])

    for appt_id in appt_ids[:2]:
        resp = client.patch(f"/api/v1/appointments/{appt_id}/cancel", headers=_auth_header(token))
        assert resp.status_code == 200

    report = client.get(f"/api/v1/reports/monthly?year={year}&month={month}", headers=_auth_header(token)).json()
    cr = report["cancellation_rate"]
    assert cr["numerator"] == 2
    assert cr["denominator"] == 4
    assert cr["value"] == pytest.approx(0.5)


# --- busiest day/hour: verified manually against known bookings -------------------------


def test_busiest_day_and_hour_match_real_bookings(business_ready):
    token = business_ready["token"]
    year, month = _this_year_month()
    # A 15-minute service so 3 non-overlapping bookings fit inside one hour
    # (00/15/30) — this service has no assigned staff, so it occupies a
    # single shared per-business resource and the DB exclusion constraint
    # refuses any overlapping booking (see booking_service's module docstring).
    short_service = _create_service(token, name="Quick Checkup", duration_minutes=15)

    busy_day = _day_in_current_month(5)
    quiet_day = _day_in_current_month(6)
    if busy_day == quiet_day:
        pytest.skip("not enough days left in the current month to isolate two distinct days")

    for i in range(3):
        cust = _create_customer(token, phone=f"+1555555000{i}")
        when = datetime.combine(busy_day, datetime.min.time()).replace(hour=9, minute=i * 15, tzinfo=ZoneInfo("UTC"))
        _book(token, short_service, uuid.UUID(cust["id"]), when)
    cust = _create_customer(token, phone="+15555550099")
    when = datetime.combine(quiet_day, datetime.min.time()).replace(hour=15, tzinfo=ZoneInfo("UTC"))
    _book(token, business_ready["service_id"], uuid.UUID(cust["id"]), when)

    report = client.get(f"/api/v1/reports/monthly?year={year}&month={month}", headers=_auth_header(token)).json()

    expected_busy_day_name = calendar.day_name[busy_day.weekday()]
    assert report["busiest_days"][0] == {"day": expected_busy_day_name, "count": 3}
    assert report["busiest_hours"][0] == {"hour": 9, "count": 3}


def test_most_requested_services_matches_real_bookings(business_ready):
    token = business_ready["token"]
    year, month = _this_year_month()
    service_popular = business_ready["service_id"]
    service_rare = _create_service(token, name="Whitening")

    for i in range(3):
        cust = _create_customer(token, phone=f"+1555666000{i}")
        when = datetime.combine(_day_in_current_month(7), datetime.min.time()).replace(hour=9 + i, tzinfo=ZoneInfo("UTC"))
        _book(token, service_popular, uuid.UUID(cust["id"]), when)
    cust = _create_customer(token, phone="+15556660099")
    when = datetime.combine(_day_in_current_month(7), datetime.min.time()).replace(hour=14, tzinfo=ZoneInfo("UTC"))
    _book(token, service_rare, uuid.UUID(cust["id"]), when)

    report = client.get(f"/api/v1/reports/monthly?year={year}&month={month}", headers=_auth_header(token)).json()
    top = report["most_requested_services"][0]
    assert top["service_name"] == "Cleaning"
    assert top["count"] == 3


# --- cross-tenant isolation --------------------------------------------------------------


def test_monthly_report_cross_tenant_isolation(two_businesses):
    token_a, token_b = two_businesses["token_a"], two_businesses["token_b"]
    _open_all_week(token_a)
    _open_all_week(token_b)
    service_a = _create_service(token_a)
    service_b = _create_service(token_b)
    year, month = _current_year_month()
    when = _future_slot(8, 10)

    cust_a = _create_customer(token_a, phone="+15557770001")
    cust_b = _create_customer(token_b, phone="+15557770002")
    appt_a = _book(token_a, service_a, uuid.UUID(cust_a["id"]), when)
    appt_b = _book(token_b, service_b, uuid.UUID(cust_b["id"]), when)

    report_a = client.get(f"/api/v1/reports/monthly?year={year}&month={month}", headers=_auth_header(token_a)).json()
    report_b = client.get(f"/api/v1/reports/monthly?year={year}&month={month}", headers=_auth_header(token_b)).json()

    assert report_a["appointments"]["requested"] == 1
    assert report_b["appointments"]["requested"] == 1
    assert appt_a["id"] != appt_b["id"]
    assert report_a["business_name"] == "Monthly Report A"
    assert report_b["business_name"] == "Monthly Report B"


# --- month boundary uses business LOCAL timezone, not UTC -------------------------------


def test_month_boundary_uses_business_local_timezone_not_utc():
    """The real timezone-bug-prone case the ticket calls out explicitly: an
    appointment at 23:00 America/New_York on the last day of August is
    03:00 UTC on September 1st. It must count toward AUGUST's report (local
    calendar), never September's."""
    with SessionLocal() as db:
        business = Business(name="TZ Boundary Co", timezone="America/New_York")
        db.add(business)
        db.flush()
        customer = Customer(business_id=business.id, name="TZ Customer")
        service = Service(business_id=business.id, name="Cleaning", price=50, duration_minutes=30)
        db.add_all([customer, service])
        db.flush()

        local_dt = datetime(2026, 8, 31, 23, 0, tzinfo=ZoneInfo("America/New_York"))
        appt = Appointment(
            business_id=business.id,
            customer_id=customer.id,
            service_id=service.id,
            scheduled_at=local_dt,
            duration_minutes=30,
            status=AppointmentStatus.CONFIRMED,
        )
        db.add(appt)
        db.commit()
        assert local_dt.astimezone(ZoneInfo("UTC")).month == 9, "sanity check: this really is September in UTC"
        business_id = business.id

    with SessionLocal() as db:
        august = monthly_report_service.generate_monthly_report(db, business_id=business_id, year=2026, month=8)
        september = monthly_report_service.generate_monthly_report(db, business_id=business_id, year=2026, month=9)

    assert august["appointments"]["scheduled_for_month"] == 1
    assert september["appointments"]["scheduled_for_month"] == 0

    with SessionLocal() as db:
        db.delete(db.get(Business, business_id))
        db.commit()


# --- RBAC ---------------------------------------------------------------------------------


def test_staff_forbidden_from_all_monthly_report_endpoints(staff_token):
    year, month = _this_year_month()
    assert (
        client.get(f"/api/v1/reports/monthly?year={year}&month={month}", headers=_auth_header(staff_token)).status_code
        == 403
    )
    assert (
        client.get(
            f"/api/v1/reports/monthly/excel?year={year}&month={month}", headers=_auth_header(staff_token)
        ).status_code
        == 403
    )
    assert (
        client.post(
            f"/api/v1/reports/monthly/send?year={year}&month={month}", headers=_auth_header(staff_token)
        ).status_code
        == 403
    )


# --- Excel export ---------------------------------------------------------------------------


def test_monthly_excel_export_has_correct_sheets_and_data(business_ready):
    token, service_id = business_ready["token"], business_ready["service_id"]
    year, month = _current_year_month()
    cust = _create_customer(token, phone="+15558880001")
    when = _future_slot(9, 13)
    _book(token, service_id, uuid.UUID(cust["id"]), when)

    resp = client.get(f"/api/v1/reports/monthly/excel?year={year}&month={month}", headers=_auth_header(token))
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    import io

    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["Summary", "Busiest Days", "Busiest Hours", "Most Requested Services"]
    summary = {row[0].value: row[1].value for row in wb["Summary"].iter_rows(min_row=2)}
    assert summary["Business"] == "Monthly Report A"
    assert summary["Appointments Requested"] == 1


# --- email delivery --------------------------------------------------------------------------


def test_send_monthly_report_email_attaches_real_xlsx_stubbed_network(business_ready, monkeypatch):
    token, service_id = business_ready["token"], business_ready["service_id"]
    year, month = _this_year_month()
    cust = _create_customer(token, phone="+15559990001")
    when = datetime.combine(_day_in_current_month(10), datetime.min.time()).replace(hour=16, tzinfo=ZoneInfo("UTC"))
    _book(token, service_id, uuid.UUID(cust["id"]), when)
    client.patch("/api/v1/business/me", json={"email": "owner-monthly@example.com"}, headers=_auth_header(token))

    captured = {}

    class _FakeEmailProvider:
        def send(self, *, to, subject, body, html_body=None, attachments=None, credentials=None):
            captured["to"] = to
            captured["attachments"] = attachments
            captured["html_body"] = html_body
            return "250 ok"

    monkeypatch.setattr(monthly_report_service, "EmailNotificationProvider", lambda: _FakeEmailProvider())

    with SessionLocal() as db:
        result = monthly_report_service.send_monthly_report_email(
            db, business_id=business_ready["business_id"], year=year, month=month
        )

    assert result["sent"] is True
    assert result["recipient"] == "owner-monthly@example.com"
    assert captured["attachments"][0][0] == f"monthly_report_{year:04d}-{month:02d}.xlsx"
    assert "Monthly Report A" in captured["html_body"]

    import io

    wb = load_workbook(io.BytesIO(captured["attachments"][0][1]))
    assert wb.sheetnames == ["Summary", "Busiest Days", "Busiest Hours", "Most Requested Services"]


def test_send_monthly_report_endpoint_owner_can_trigger_it(business_ready, monkeypatch):
    year, month = _this_year_month()

    class _FakeEmailProvider:
        def send(self, *, to, subject, body, html_body=None, attachments=None, credentials=None):
            return "250 ok"

    monkeypatch.setattr(monthly_report_service, "EmailNotificationProvider", lambda: _FakeEmailProvider())
    client.patch(
        "/api/v1/business/me", json={"email": "owner-monthly2@example.com"}, headers=_auth_header(business_ready["token"])
    )

    resp = client.post(
        f"/api/v1/reports/monthly/send?year={year}&month={month}", headers=_auth_header(business_ready["token"])
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["sent"] is True
