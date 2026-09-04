"""Phase 16 — daily reports + Excel export (app/services/reporting/,
app/api/routes/reports.py).

Real DB and real HTTP throughout — appointments/cancellations/reschedules/
customers are created through the actual booking_service/customer_service
code paths (never inserted as pre-shaped fixtures) so the report is proven
against genuinely-produced data, not a fabricated snapshot. Only the SMTP
network call in the report-email tests is stubbed (same discipline as
test_notifications.py / test_sms_notifications.py).
"""

import io
import uuid
from datetime import date, datetime, timedelta
from zoneinfo import ZoneInfo

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.core.security import create_access_token
from app.db.database import SessionLocal
from app.db.models.appointment import Appointment
from app.db.models.audit_log import AuditLog
from app.db.models.business import Business, BusinessUser, BusinessUserRole
from app.main import app
from app.services.notifications.base import NotificationDeliveryError
from app.services.reporting import report_service

client = TestClient(app)


def _unique_email(label: str) -> str:
    return f"{label}-{uuid.uuid4().hex[:10]}@example.com"


def _auth_header(token: str) -> dict:
    return {"Authorization": f"Bearer {token}"}


def _next_weekday(target_weekday: int) -> date:
    today = date.today()
    days_ahead = (target_weekday - today.weekday()) % 7
    days_ahead = days_ahead or 7
    return today + timedelta(days=days_ahead)


@pytest.fixture
def two_businesses():
    email_a = _unique_email("rpt-a-owner")
    email_b = _unique_email("rpt-b-owner")
    resp_a = client.post(
        "/api/v1/auth/register",
        json={"business_name": "Report Test A", "timezone": "UTC", "email": email_a, "password": "correcthorse1"},
    )
    resp_b = client.post(
        "/api/v1/auth/register",
        json={"business_name": "Report Test B", "timezone": "UTC", "email": email_b, "password": "correcthorse1"},
    )
    assert resp_a.status_code == 201, resp_a.text
    assert resp_b.status_code == 201, resp_b.text
    login_a = client.post("/api/v1/auth/login", json={"email": email_a, "password": "correcthorse1"})
    login_b = client.post("/api/v1/auth/login", json={"email": email_b, "password": "correcthorse1"})

    data = {
        "business_id_a": resp_a.json()["business_id"],
        "business_id_b": resp_b.json()["business_id"],
        "token_a": login_a.json()["access_token"],
        "token_b": login_b.json()["access_token"],
    }
    yield data

    with SessionLocal() as db:
        for business_id in (data["business_id_a"], data["business_id_b"]):
            business = db.get(Business, uuid.UUID(business_id))
            if business is not None:
                db.delete(business)
        db.commit()


@pytest.fixture
def staff_token(two_businesses):
    business_id_a = uuid.UUID(two_businesses["business_id_a"])
    with SessionLocal() as db:
        staff_user = BusinessUser(
            business_id=business_id_a,
            email=_unique_email("rpt-staff"),
            hashed_password="not-used-in-this-test",
            role=BusinessUserRole.STAFF,
        )
        db.add(staff_user)
        db.commit()
        db.refresh(staff_user)
        return create_access_token(user_id=staff_user.id, business_id=business_id_a, role=staff_user.role.value)


def _open_all_week(token: str):
    days = [{"day_of_week": d, "open_time": "00:00:00", "close_time": "23:45:00"} for d in range(7)]
    assert client.put("/api/v1/business/hours", json={"days": days}, headers=_auth_header(token)).status_code == 200


def _create_service(token: str) -> uuid.UUID:
    resp = client.post(
        "/api/v1/services", json={"name": "Cleaning", "price": "50.00", "duration_minutes": 30}, headers=_auth_header(token)
    )
    assert resp.status_code == 201, resp.text
    return uuid.UUID(resp.json()["id"])


def _create_customer(token: str, **kwargs) -> dict:
    payload = {"name": "Report Customer", **kwargs}
    resp = client.post("/api/v1/customers", json=payload, headers=_auth_header(token))
    assert resp.status_code == 201, resp.text
    return resp.json()


def _book(token: str, service_id: uuid.UUID, customer_id: uuid.UUID, when: datetime) -> dict:
    resp = client.post(
        "/api/v1/appointments",
        json={"customer_id": str(customer_id), "service_id": str(service_id), "scheduled_at": when.isoformat()},
        headers=_auth_header(token),
    )
    assert resp.status_code == 201, resp.text
    return resp.json()


@pytest.fixture
def business_ready(two_businesses):
    token = two_businesses["token_a"]
    _open_all_week(token)
    service_id = _create_service(token)
    return {"token": token, "business_id": uuid.UUID(two_businesses["business_id_a"]), "service_id": service_id}


# --- appointments-scheduled section ------------------------------------------------


def test_report_appointments_section_matches_real_bookings(business_ready):
    day = _next_weekday(0)
    customer = _create_customer(business_ready["token"], phone="+15550000001")
    when = datetime.combine(day, datetime.min.time()).replace(hour=10, tzinfo=ZoneInfo("UTC"))
    appt = _book(business_ready["token"], business_ready["service_id"], uuid.UUID(customer["id"]), when)

    resp = client.get(f"/api/v1/reports/daily?date={day.isoformat()}", headers=_auth_header(business_ready["token"]))
    assert resp.status_code == 200, resp.text
    report = resp.json()

    assert report["report_date"] == day.isoformat()
    ids = [a["id"] for a in report["appointments"]]
    assert appt["id"] in ids
    row = next(a for a in report["appointments"] if a["id"] == appt["id"])
    assert row["customer_name"] == "Report Customer"
    assert row["service_name"] == "Cleaning"
    assert row["status"] == "confirmed"
    assert report["summary"]["appointments_scheduled"] == 1
    assert report["summary"]["appointments_by_status"] == {"confirmed": 1}


def test_report_zero_activity_day_is_honest_empty(business_ready):
    far_future = date.today() + timedelta(days=365)
    resp = client.get(
        f"/api/v1/reports/daily?date={far_future.isoformat()}", headers=_auth_header(business_ready["token"])
    )
    assert resp.status_code == 200, resp.text
    report = resp.json()
    assert report["appointments"] == []
    assert report["cancellations"] == []
    assert report["reschedules"] == []
    assert report["new_leads"] == []
    assert report["summary"] == {
        "appointments_scheduled": 0,
        "appointments_by_status": {},
        "cancellations": 0,
        "reschedules": 0,
        "new_leads": 0,
        "human_review_open_count": 0,
    }


# --- cancellations section (keyed by real cancellation event date) -----------------


def test_report_cancellations_section_reflects_real_cancellation_event(business_ready):
    day = _next_weekday(1)
    customer = _create_customer(business_ready["token"], phone="+15550000002")
    when = datetime.combine(day, datetime.min.time()).replace(hour=11, tzinfo=ZoneInfo("UTC"))
    appt = _book(business_ready["token"], business_ready["service_id"], uuid.UUID(customer["id"]), when)

    cancel_resp = client.patch(
        f"/api/v1/appointments/{appt['id']}/cancel", headers=_auth_header(business_ready["token"])
    )
    assert cancel_resp.status_code == 200, cancel_resp.text

    with SessionLocal() as db:
        row = db.get(Appointment, uuid.UUID(appt["id"]))
        cancelled_at_date = row.updated_at.date()

    resp = client.get(
        f"/api/v1/reports/daily?date={cancelled_at_date.isoformat()}", headers=_auth_header(business_ready["token"])
    )
    report = resp.json()
    ids = [c["id"] for c in report["cancellations"]]
    assert appt["id"] in ids
    row = next(c for c in report["cancellations"] if c["id"] == appt["id"])
    assert row["originally_scheduled_at"] == when.isoformat()
    assert row["customer_name"] == "Report Customer"
    assert report["summary"]["cancellations"] == 1


# --- reschedules section (real AuditLog trail) --------------------------------------


def test_report_reschedules_section_shows_real_old_and_new_time(business_ready):
    day = _next_weekday(2)
    customer = _create_customer(business_ready["token"], phone="+15550000003")
    old_time = datetime.combine(day, datetime.min.time()).replace(hour=9, tzinfo=ZoneInfo("UTC"))
    new_time = datetime.combine(day, datetime.min.time()).replace(hour=13, tzinfo=ZoneInfo("UTC"))
    appt = _book(business_ready["token"], business_ready["service_id"], uuid.UUID(customer["id"]), old_time)

    resched = client.patch(
        f"/api/v1/appointments/{appt['id']}/reschedule",
        json={"scheduled_at": new_time.isoformat()},
        headers=_auth_header(business_ready["token"]),
    )
    assert resched.status_code == 200, resched.text

    with SessionLocal() as db:
        log = (
            db.query(AuditLog)
            .filter(AuditLog.resource_id == appt["id"], AuditLog.action == "appointment_rescheduled")
            .one()
        )
        changed_date = log.created_at.date()

    resp = client.get(
        f"/api/v1/reports/daily?date={changed_date.isoformat()}", headers=_auth_header(business_ready["token"])
    )
    report = resp.json()
    ids = [r["id"] for r in report["reschedules"]]
    assert appt["id"] in ids
    row = next(r for r in report["reschedules"] if r["id"] == appt["id"])
    assert row["old_scheduled_at"] == old_time.isoformat()
    assert row["new_scheduled_at"] == new_time.isoformat()
    assert report["summary"]["reschedules"] == 1


def test_report_reschedules_section_handles_two_reschedules_same_day(business_ready):
    """Proves the AuditLog-chain reconstruction, not just the single-hop case:
    the FIRST reschedule's "new time" must equal the SECOND reschedule's "old
    time" (both derived from real audit rows, never guessed), and the SECOND
    reschedule's "new time" must match the appointment's real current
    scheduled_at."""
    day = _next_weekday(3)
    customer = _create_customer(business_ready["token"], phone="+15550000004")
    t0 = datetime.combine(day, datetime.min.time()).replace(hour=8, tzinfo=ZoneInfo("UTC"))
    t1 = datetime.combine(day, datetime.min.time()).replace(hour=12, tzinfo=ZoneInfo("UTC"))
    t2 = datetime.combine(day, datetime.min.time()).replace(hour=16, tzinfo=ZoneInfo("UTC"))
    appt = _book(business_ready["token"], business_ready["service_id"], uuid.UUID(customer["id"]), t0)

    r1 = client.patch(
        f"/api/v1/appointments/{appt['id']}/reschedule",
        json={"scheduled_at": t1.isoformat()},
        headers=_auth_header(business_ready["token"]),
    )
    assert r1.status_code == 200, r1.text
    r2 = client.patch(
        f"/api/v1/appointments/{appt['id']}/reschedule",
        json={"scheduled_at": t2.isoformat()},
        headers=_auth_header(business_ready["token"]),
    )
    assert r2.status_code == 200, r2.text

    with SessionLocal() as db:
        logs = (
            db.query(AuditLog)
            .filter(AuditLog.resource_id == appt["id"], AuditLog.action == "appointment_rescheduled")
            .order_by(AuditLog.created_at)
            .all()
        )
        assert len(logs) == 2
        changed_date = logs[0].created_at.date()

    resp = client.get(
        f"/api/v1/reports/daily?date={changed_date.isoformat()}", headers=_auth_header(business_ready["token"])
    )
    report = resp.json()
    rows = [r for r in report["reschedules"] if r["id"] == appt["id"]]
    assert len(rows) == 2, "both reschedule events on the same day must both be reported, not collapsed"
    rows.sort(key=lambda r: r["changed_at"])
    assert rows[0]["old_scheduled_at"] == t0.isoformat()
    assert rows[0]["new_scheduled_at"] == t1.isoformat()
    assert rows[1]["old_scheduled_at"] == t1.isoformat()
    assert rows[1]["new_scheduled_at"] == t2.isoformat(), "second hop's new time must match the real current appointment row"


# --- new leads section ---------------------------------------------------------------


def test_report_new_leads_section_reflects_real_customer_creation(business_ready):
    customer = _create_customer(business_ready["token"], phone="+15550000005", email="lead@example.com")
    with SessionLocal() as db:
        from app.db.models.customer import Customer

        row = db.get(Customer, uuid.UUID(customer["id"]))
        created_date = row.created_at.date()

    resp = client.get(
        f"/api/v1/reports/daily?date={created_date.isoformat()}", headers=_auth_header(business_ready["token"])
    )
    report = resp.json()
    ids = [c["id"] for c in report["new_leads"]]
    assert customer["id"] in ids
    row = next(c for c in report["new_leads"] if c["id"] == customer["id"])
    assert row["name"] == "Report Customer"
    assert row["phone"] == "+15550000005"
    assert row["email"] == "lead@example.com"


# --- human review section: real since Phase 19's HumanHandoff producer -----------------


def test_report_human_review_section_is_zero_but_real_with_no_handoffs(business_ready):
    today = date.today()
    resp = client.get(f"/api/v1/reports/daily?date={today.isoformat()}", headers=_auth_header(business_ready["token"]))
    report = resp.json()
    assert report["human_review"]["implemented"] is True
    assert report["human_review"]["count"] == 0


def test_report_human_review_section_reflects_a_real_open_handoff(business_ready):
    """Phase 19 gave HumanHandoff its first real producer — this is no longer
    always-empty (Phase 16's own flagged gap)."""
    from app.db.models.conversation import Conversation
    from app.schemas.conversation import ConversationIntent
    from app.services import handoff_service

    customer = _create_customer(business_ready["token"], phone="+15550000099")
    with SessionLocal() as db:
        conversation = Conversation(
            business_id=business_ready["business_id"],
            customer_id=uuid.UUID(customer["id"]),
            channel="sms",
            status="open",
        )
        db.add(conversation)
        db.commit()
        db.refresh(conversation)
        handoff_service.maybe_create_handoff(
            db,
            business_id=business_ready["business_id"],
            conversation_id=conversation.id,
            intent=ConversationIntent.COMPLAINT,
            best_similarity=None,
        )

    today = date.today()
    resp = client.get(f"/api/v1/reports/daily?date={today.isoformat()}", headers=_auth_header(business_ready["token"]))
    report = resp.json()
    assert report["human_review"]["implemented"] is True
    assert report["human_review"]["count"] == 1
    assert report["summary"]["human_review_open_count"] == 1


# --- cross-tenant isolation ------------------------------------------------------------


def test_report_cross_tenant_isolation_with_overlapping_dates(two_businesses):
    token_a, token_b = two_businesses["token_a"], two_businesses["token_b"]
    _open_all_week(token_a)
    _open_all_week(token_b)
    service_a = _create_service(token_a)
    service_b = _create_service(token_b)

    day = _next_weekday(4)
    when = datetime.combine(day, datetime.min.time()).replace(hour=10, tzinfo=ZoneInfo("UTC"))

    customer_a = _create_customer(token_a, phone="+15550000006")
    customer_b = _create_customer(token_b, phone="+15550000007")
    appt_a = _book(token_a, service_a, uuid.UUID(customer_a["id"]), when)
    appt_b = _book(token_b, service_b, uuid.UUID(customer_b["id"]), when)

    report_a = client.get(f"/api/v1/reports/daily?date={day.isoformat()}", headers=_auth_header(token_a)).json()
    report_b = client.get(f"/api/v1/reports/daily?date={day.isoformat()}", headers=_auth_header(token_b)).json()

    ids_a = [a["id"] for a in report_a["appointments"]]
    ids_b = [a["id"] for a in report_b["appointments"]]
    assert appt_a["id"] in ids_a
    assert appt_a["id"] not in ids_b
    assert appt_b["id"] in ids_b
    assert appt_b["id"] not in ids_a
    assert report_a["business_name"] == "Report Test A"
    assert report_b["business_name"] == "Report Test B"


# --- RBAC --------------------------------------------------------------------------


def test_staff_forbidden_from_all_three_report_endpoints(staff_token):
    today = date.today().isoformat()
    assert client.get(f"/api/v1/reports/daily?date={today}", headers=_auth_header(staff_token)).status_code == 403
    assert (
        client.get(f"/api/v1/reports/daily/excel?date={today}", headers=_auth_header(staff_token)).status_code == 403
    )
    assert (
        client.post(f"/api/v1/reports/daily/send?date={today}", headers=_auth_header(staff_token)).status_code == 403
    )


# --- Excel export --------------------------------------------------------------------


def test_excel_export_produces_a_real_readable_xlsx_with_correct_sheets_and_data(business_ready):
    day = _next_weekday(5)
    customer = _create_customer(business_ready["token"], phone="+15550000008")
    when = datetime.combine(day, datetime.min.time()).replace(hour=14, tzinfo=ZoneInfo("UTC"))
    appt = _book(business_ready["token"], business_ready["service_id"], uuid.UUID(customer["id"]), when)

    resp = client.get(
        f"/api/v1/reports/daily/excel?date={day.isoformat()}", headers=_auth_header(business_ready["token"])
    )
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert f"daily_report_{day.isoformat()}.xlsx" in resp.headers["content-disposition"]
    assert len(resp.content) > 1000, "a real .xlsx (zip container) is never a trivially small byte string"

    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["Appointments", "Cancellations", "Reschedules", "New Leads", "Summary"]

    ws = wb["Appointments"]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert header == ["Time", "Customer", "Service", "Staff", "Status", "Booking ID"]
    data_row = [c.value for c in next(ws.iter_rows(min_row=2, max_row=2))]
    assert data_row[1] == "Report Customer"
    assert data_row[2] == "Cleaning"
    assert data_row[4] == "confirmed"
    assert data_row[5] == appt["id"]

    ws = wb["Summary"]
    summary_rows = {row[0].value: row[1].value for row in ws.iter_rows(min_row=2)}
    assert summary_rows["Business"] == "Report Test A"
    assert summary_rows["Appointments Scheduled"] == 1


# --- email delivery ------------------------------------------------------------------


def test_send_daily_report_email_no_recipient_configured(business_ready):
    """business.email is never set at registration (only the owner login
    email is) — the honest default state, proving the "no recipient" branch
    never crashes and reports a clear reason."""
    today = date.today()
    with SessionLocal() as db:
        result = report_service.send_daily_report_email(db, business_id=business_ready["business_id"], report_date=today)
    assert result["sent"] is False
    assert "no report email" in result["reason"].lower() or "email" in result["reason"].lower()


def test_send_daily_report_email_attaches_a_real_xlsx_stubbed_network(business_ready, monkeypatch):
    day = _next_weekday(6)
    customer = _create_customer(business_ready["token"], phone="+15550000009")
    when = datetime.combine(day, datetime.min.time()).replace(hour=15, tzinfo=ZoneInfo("UTC"))
    _book(business_ready["token"], business_ready["service_id"], uuid.UUID(customer["id"]), when)

    set_email = client.patch(
        "/api/v1/business/me", json={"email": "owner-report@example.com"}, headers=_auth_header(business_ready["token"])
    )
    assert set_email.status_code == 200, set_email.text

    captured = {}

    class _FakeEmailProvider:
        def send(self, *, to, subject, body, html_body=None, attachments=None):
            captured["to"] = to
            captured["subject"] = subject
            captured["attachments"] = attachments
            return "250 message accepted for delivery"

    monkeypatch.setattr(report_service, "EmailNotificationProvider", lambda: _FakeEmailProvider())

    with SessionLocal() as db:
        result = report_service.send_daily_report_email(db, business_id=business_ready["business_id"], report_date=day)

    assert result["sent"] is True
    assert result["recipient"] == "owner-report@example.com"
    assert result["attachment_size_bytes"] > 1000
    assert captured["to"] == "owner-report@example.com"
    filename, content, mime_type = captured["attachments"][0]
    assert filename == f"daily_report_{day.isoformat()}.xlsx"
    assert mime_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    # the attached bytes are a real, independently-readable workbook, not a placeholder
    wb = load_workbook(io.BytesIO(content))
    assert wb.sheetnames == ["Appointments", "Cancellations", "Reschedules", "New Leads", "Summary"]


def test_send_daily_report_email_provider_failure_never_crashes(business_ready, monkeypatch):
    client.patch(
        "/api/v1/business/me", json={"email": "owner-report@example.com"}, headers=_auth_header(business_ready["token"])
    )

    class _FailingProvider:
        def send(self, *, to, subject, body, html_body=None, attachments=None):
            raise NotificationDeliveryError("smtp down", transient=True)

    monkeypatch.setattr(report_service, "EmailNotificationProvider", lambda: _FailingProvider())

    with SessionLocal() as db:
        result = report_service.send_daily_report_email(
            db, business_id=business_ready["business_id"], report_date=date.today()
        )
    assert result["sent"] is False
    assert result["transient"] is True


def test_send_daily_report_endpoint_owner_can_trigger_it(business_ready, monkeypatch):
    class _FakeEmailProvider:
        def send(self, *, to, subject, body, html_body=None, attachments=None):
            return "250 ok"

    monkeypatch.setattr(report_service, "EmailNotificationProvider", lambda: _FakeEmailProvider())
    client.patch(
        "/api/v1/business/me", json={"email": "owner-report@example.com"}, headers=_auth_header(business_ready["token"])
    )

    resp = client.post(
        f"/api/v1/reports/daily/send?date={date.today().isoformat()}", headers=_auth_header(business_ready["token"])
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["sent"] is True
